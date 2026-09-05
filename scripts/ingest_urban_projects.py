"""서울 도시계획사업(UQ120) → 상권 겹침 + 자치구 요약. 프로파일 FC-51·52.

원천: 서울 도시계획포털 UQ120_도시계획사업 (UPIS_C_UQ120.shp, EPSG:5174, dbf cp949)
      + 서울플랜+코드정의표(대민용).xlsx (사업유형·추진단계 코드)
  사업유형: 정비사업(재개발·재건축·모아타운 등)·역세권사업·도시개발·재정비촉진 등 폴리곤
  추진단계: 대상지선정 → 구역지정 → 조합설립 → 사업시행인가 → 관리처분 → 착공 → 준공

산출(집계본만 커밋):
  data/도시계획사업/도시계획사업_상권겹침.csv   — 상권 × 사업 (겹침면적·비율)
  data/도시계획사업/도시계획사업_자치구요약.csv — 자치구 × 유형 × 단계구분 건수

실행: .venv/bin/python3 scripts/ingest_urban_projects.py [--src UQ120_도시계획사업]
"""
from __future__ import annotations
import csv, os, sys, collections
import shapefile
import openpyxl
from shapely.geometry import shape
from shapely.ops import transform as shp_transform
from shapely.strtree import STRtree
from pyproj import Transformer

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from _raw import raw
SRC = os.path.expanduser(sys.argv[sys.argv.index("--src") + 1]) if "--src" in sys.argv \
      else raw("도시계획사업", "UQ120_도시계획사업")
OUT_DIR = os.path.join(ROOT, "data", "도시계획사업")
AREA_SHP = os.path.join(ROOT, "data", "영역", "상권", "서울시 상권분석서비스(영역-상권)")

MIN_OVERLAP_M2 = 100.0

# 추진단계 한글명 → 진행도 구분 (초기 < 조합 < 인가 < 착공 < 완료)
def phase_bucket(name: str) -> str:
    if any(k in name for k in ("준공", "입주", "사용승인", "사업완료")):
        return "완료"
    if "착공" in name:
        return "착공"
    if any(k in name for k in ("사업시행", "관리처분", "사업계획승인", "실시계획", "지구계획", "사업계획인가", "사업시행계획")):
        return "인가"
    if any(k in name for k in ("조합설립", "추진위", "주민합의체", "추진계획승인")):
        return "조합"
    return "초기"


def load_codes():
    p = os.path.join(SRC, "서울플랜+코드정의표(대민용).xlsx")
    wb = openpyxl.load_workbook(p, read_only=True)
    sclas = {}
    lclas = {}
    ws = wb["01. 사업유형_코드"]
    cur_l = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        lc = row[1]
        if lc:
            cur_l = lc
        if row[4]:
            sclas[row[4]] = row[3]
            lclas[row[4]] = row[0] or lclas_prev if (row[0]) else None
    # 대분류: 첫 열 병합값 채우기
    ws = wb["01. 사업유형_코드"]
    big = None
    lbl = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            big = row[0]
        if row[4]:
            lbl[row[4]] = big
    propel = {}
    ws = wb["02. 추진단계_코드"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[2] and row[3]:
            propel[row[2]] = row[3]
    wb.close()
    return sclas, lbl, propel


def load_areas():
    sf = shapefile.Reader(AREA_SHP, encoding="utf-8")
    fields = [f[0] for f in sf.fields if f[0] != "DeletionFlag"]
    geoms, codes, names, gus = [], [], [], []
    gu_name_by_code = {}
    for sr in sf.iterShapeRecords():
        rec = dict(zip(fields, sr.record))
        geoms.append(shape(sr.shape.__geo_interface__))
        codes.append(str(rec["TRDAR_CD"]))
        names.append(rec["TRDAR_CD_N"])
        gus.append(rec.get("SIGNGU_CD_", ""))
        gu_name_by_code[str(rec.get("SIGNGU_CD", ""))] = rec.get("SIGNGU_CD_", "")
    return STRtree(geoms), geoms, codes, names, gus, gu_name_by_code


def main() -> int:
    shp = os.path.join(SRC, "shp파일", "UPIS_C_UQ120")
    if not os.path.isfile(shp + ".shp"):
        print(f"FAIL: {shp}.shp 없음")
        return 1
    sclas, lbl, propel = load_codes()
    tree, ageoms, acodes, anames, agus, gu_name = load_areas()
    to5181 = Transformer.from_crs("EPSG:5174", "EPSG:5181", always_xy=True).transform

    sf = shapefile.Reader(shp, encoding="cp949")
    fields = [f[0] for f in sf.fields if f[0] != "DeletionFlag"]
    os.makedirs(OUT_DIR, exist_ok=True)

    overlap_rows = []
    gu_summary = collections.Counter()   # (자치구, 대분류, 단계구분) -> n
    n_proj = n_no_area = 0
    typ_dist = collections.Counter()

    for sr in sf.iterShapeRecords():
        rec = dict(zip(fields, sr.record))
        n_proj += 1
        sc = str(rec.get("SCLAS_CL", "")).strip()
        typ = sclas.get(sc, sc or "미상")
        big = lbl.get(sc, "미상")
        stage = propel.get(str(rec.get("PROPEL_CD", "")).strip(),
                           str(rec.get("PROPEL_CD", "")).strip() or "미상")
        bucket = phase_bucket(stage)
        gu_cd = str(rec.get("SIGNGU_SE", "")).strip()
        gu = gu_name.get(gu_cd, gu_cd)
        nm = str(rec.get("DGM_NM", "")).strip()
        created = str(rec.get("CREATE_DAT", "")).strip()
        typ_dist[big] += 1

        try:
            g5174 = shape(sr.shape.__geo_interface__)
            g = shp_transform(lambda x, y, z=None: to5181(x, y), g5174)
        except Exception:
            continue
        if not g.is_valid:
            g = g.buffer(0)

        hit_any = False
        for idx in tree.query(g):
            inter = ageoms[idx].intersection(g).area
            if inter < MIN_OVERLAP_M2:
                continue
            hit_any = True
            overlap_rows.append([
                acodes[idx], anames[idx], agus[idx] or gu, big, typ, stage, bucket,
                nm, round(g.area), round(inter),
                round(inter / ageoms[idx].area * 100, 2) if ageoms[idx].area else "",
                created,
            ])
        if not hit_any:
            n_no_area += 1
        gu_summary[(gu_cd, big, bucket)] += 1

    ov_path = os.path.join(OUT_DIR, "도시계획사업_상권겹침.csv")
    with open(ov_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["상권_코드", "상권_코드_명", "자치구", "사업유형_대분류", "사업유형",
                    "추진단계", "추진단계_구분", "사업장명", "사업_면적", "겹침_면적",
                    "겹침_상권비율", "생성일"])
        w.writerows(sorted(overlap_rows, key=lambda r: (r[0], r[3])))

    su_path = os.path.join(OUT_DIR, "도시계획사업_자치구요약.csv")
    with open(su_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["자치구_코드", "자치구", "사업유형_대분류", "추진단계_구분", "건수"])
        for (gu_cd, big, bk), n in sorted(gu_summary.items()):
            w.writerow([gu_cd, gu_name.get(gu_cd, ""), big, bk, n])

    print(f"도시계획사업 폴리곤 {n_proj:,}건  (좌표계 EPSG:5174→5181)")
    print(f"사업유형 대분류: {dict(typ_dist.most_common())}")
    print(f"상권 겹침 레코드 {len(overlap_rows):,}  /  어느 상권에도 안 겹침 {n_no_area:,}")
    parea = {r[0] for r in overlap_rows}
    print(f"겹친 상권 {len(parea)}/1650")
    print(f"\n→ {os.path.relpath(ov_path, ROOT)}")
    print(f"→ {os.path.relpath(su_path, ROOT)}")
    # 샘플
    print("\n샘플 (송파구):")
    for r in overlap_rows:
        if r[2] == "송파구" and r[6] in ("인가", "착공"):
            print(f"  {r[1]}({r[0]}) ← {r[4]} / {r[5]} / {r[7]}  겹침 {r[10]}%")
    return 0


if __name__ == "__main__":
    sys.exit(main())
