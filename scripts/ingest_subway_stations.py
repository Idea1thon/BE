"""전체 도시철도역사정보 → 서울 역사 목록 + 상권별 역세권 접근성.

원천: 국가철도공단_전체 도시철도역사정보 (공공데이터포털 15013205)
      `전체_도시철도역사정보_YYYYMMDD.xlsx`, 1개 시트, 15열, 좌표계 WGS84(역위도·역경도).
      전국 1,099역. 노선명 표기 불일치("9호선"/"서울 도시철도 9호선"/"수도권 도시철도 9호선").
      개통 스냅샷 — 개통일 시계열 아님.

산출(집계본만 커밋, 원천 xlsx 는 .gitignore):
  data/도시철도역사/역사정보_서울.csv    — 역-노선 단위(서울 행정구역 + 인접 경기역), 상권·행정동 배정
  data/도시철도역사/상권_역세권.csv       — 상권 1,650 단위: 상권내 역 수·최근접 역·거리·환승 인접·인접역 목록(json)

실행: .venv/bin/python3 scripts/ingest_subway_stations.py [--src FILE]
"""
from __future__ import annotations
import csv, glob, json, os, re, sys

import openpyxl
import shapefile
from shapely.geometry import shape, Point
from shapely.strtree import STRtree
from pyproj import Transformer

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from _raw import raw

if "--src" in sys.argv:
    SRC = os.path.expanduser(sys.argv[sys.argv.index("--src") + 1])
else:
    cands = sorted(glob.glob(raw("도시철도역사", "전체_도시철도역사정보_*.xlsx")))
    SRC = cands[-1] if cands else raw("도시철도역사", "전체_도시철도역사정보_20260630.xlsx")

OUT_DIR = os.path.join(ROOT, "data", "도시철도역사")
TRDAR_SHP = os.path.join(ROOT, "data", "영역", "상권", "서울시 상권분석서비스(영역-상권)")
DONG_SHP = os.path.join(ROOT, "data", "영역", "행정동", "서울시 상권분석서비스(영역-행정동)")

NEAR_M = 500     # nearby_anchors 후보로 노출할 상권↔역 최대 거리(m)
YEOKSEGWON_M = 250   # "역세권" 관용 기준(역 중심 250m)
MARGIN_M = 1200   # 서울 행정구역 밖이어도 이 거리 안이면 인접역으로 보존

_PREFIX = re.compile(
    r"^(서울 |수도권 |인천 )?(도시철도 |광역철도 |광역급행철도 |경량도시철도 |지하철 )+"
)
_LINE_ALIAS = {
    "인천국제공항선": "공항철도", "공항철도1호선": "공항철도",
    "경부선": "1호선", "경인선": "1호선", "경원선": "1호선", "경부고속선": "1호선",
    "장항선": "1호선",
    "수인선": "수인분당선", "분당선": "수인분당선",
    "안산과천선": "4호선", "과천선": "4호선", "안산선": "4호선",
    "일산선": "3호선",
}


def norm_line(name: str) -> str:
    n = re.sub(r"\s+", " ", (name or "").strip())
    n = _PREFIX.sub("", n)
    return _LINE_ALIAS.get(n, n)


def load_shp(path, code_field, name_field):
    sf = shapefile.Reader(path, encoding="utf-8")
    fields = [f[0] for f in sf.fields if f[0] != "DeletionFlag"]
    geoms, recs = [], []
    for sr in sf.iterShapeRecords():
        rec = dict(zip(fields, sr.record))
        geoms.append(shape(sr.shape.__geo_interface__))
        recs.append((str(rec[code_field]), rec[name_field], rec))
    return STRtree(geoms), geoms, recs


def pip(tree, geoms, recs, pt):
    for idx in tree.query(pt):
        if geoms[idx].covers(pt):
            return recs[idx]
    return None


def main() -> int:
    if not os.path.isfile(SRC):
        print(f"FAIL: {SRC} 없음")
        return 1
    os.makedirs(OUT_DIR, exist_ok=True)

    print(f"원천: {SRC}")
    wb = openpyxl.load_workbook(SRC, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() for c in next(rows)]
    H = {c: i for i, c in enumerate(header)}
    need = ["역번호", "역사명", "노선명", "환승역구분", "역위도", "역경도",
            "운영기관명", "역사도로명주소", "데이터기준일자"]
    for c in need:
        if c not in H:
            print(f"FAIL: 컬럼 '{c}' 없음 — {header}")
            return 1

    tf = Transformer.from_crs("EPSG:4326", "EPSG:5181", always_xy=True)
    print("상권·행정동 폴리곤 로드...")
    d_tree, d_geoms, d_recs = load_shp(DONG_SHP, "ADSTRD_CD", "ADSTRD_NM")
    t_tree, t_geoms, t_recs = load_shp(TRDAR_SHP, "TRDAR_CD", "TRDAR_CD_N")

    as_of = None
    raw_n = seoul_n = margin_n = drop_n = 0
    # 역-노선 레코드
    station_rows = []
    # 상권별 근접역: code -> list[(dist, 역사명, 노선, 환승bool, 내부bool)]
    near = {code: [] for code, _, _ in t_recs}

    for r in rows:
        if not r or r[H["역번호"]] is None:
            continue
        raw_n += 1
        try:
            lat = float(r[H["역위도"]]); lon = float(r[H["역경도"]])
        except (TypeError, ValueError):
            continue
        as_of = as_of or (str(r[H["데이터기준일자"]])[:10] if r[H["데이터기준일자"]] else "")
        x, y = tf.transform(lon, lat)
        pt = Point(x, y)

        dong = pip(d_tree, d_geoms, d_recs, pt)
        addr = str(r[H["역사도로명주소"]] or "")

        # 서울 행정구역(행정동 PIP) 또는 상권 근처(MARGIN) 만 보존
        if dong is None:
            nidx = t_tree.nearest(pt)
            if t_geoms[nidx].distance(pt) > MARGIN_M:
                drop_n += 1
                continue
            margin_n += 1
        else:
            seoul_n += 1

        line = norm_line(str(r[H["노선명"]]))
        is_transfer = "환승" in str(r[H["환승역구분"]] or "")
        trdar = pip(t_tree, t_geoms, t_recs, pt)
        sig = (dong[2]["ADSTRD_CD"][:5] if dong else "")

        station_rows.append({
            "역번호": r[H["역번호"]],
            "역사명": str(r[H["역사명"]]).strip(),
            "노선명_원본": str(r[H["노선명"]]).strip(),
            "노선명": line,
            "환승역": "Y" if is_transfer else "N",
            "위도": round(lat, 7), "경도": round(lon, 7),
            "X_5181": round(x, 2), "Y_5181": round(y, 2),
            "행정구역_서울": "Y" if dong is not None else "N",
            "자치구코드": sig,
            "행정동_코드": dong[0] if dong else "",
            "행정동명": dong[1] if dong else "",
            "상권_코드": trdar[0] if trdar else "",
            "상권명": trdar[1] if trdar else "",
            "역사도로명주소": addr,
            "운영기관명": str(r[H["운영기관명"]] or "").strip(),
            "데이터기준일자": str(r[H["데이터기준일자"]])[:10] if r[H["데이터기준일자"]] else "",
        })

        # 상권별 근접역 (역 1개가 여러 상권에 근접 가능) — buffer 로 후보 폴리곤만
        buf = pt.buffer(NEAR_M)
        for gi in t_tree.query(buf):
            d = t_geoms[gi].distance(pt)
            if d <= NEAR_M:
                near[t_recs[gi][0]].append(
                    (round(d, 1), str(r[H["역사명"]]).strip(), line, is_transfer, d == 0.0))

    wb.close()

    # ---- 역사정보_서울.csv ----
    st_path = os.path.join(OUT_DIR, "역사정보_서울.csv")
    cols = ["역번호", "역사명", "노선명", "노선명_원본", "환승역", "위도", "경도",
            "X_5181", "Y_5181", "행정구역_서울", "자치구코드", "행정동_코드", "행정동명",
            "상권_코드", "상권명", "역사도로명주소", "운영기관명", "데이터기준일자"]
    station_rows.sort(key=lambda s: (s["노선명"], str(s["역번호"])))
    with open(st_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(station_rows)

    # ---- 상권_역세권.csv ----
    sg_path = os.path.join(OUT_DIR, "상권_역세권.csv")
    n_with = 0
    with open(sg_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["상권_코드", "상권명", "자치구", "상권내_역_수", "최근접_역명",
                    "최근접_노선", "최근접_역_거리_m", "역세권_250m", "환승역_인접_500m",
                    "인접역_목록", "데이터기준일자"])
        for code, name, rec in t_recs:
            lst = sorted(near[code])
            inside = sum(1 for d, *_ in lst if d == 0.0)
            if lst:
                n_with += 1
                bd, bn, bl, *_ = lst[0]
                transfer_near = any(t for _, _, _, t, _ in lst)
                anchors = [{"name": n_, "line": l_, "distance_m": d_,
                            "transfer": tr_} for d_, n_, l_, tr_, _ in lst[:6]]
            else:
                bd = bn = bl = ""
                transfer_near = False
                anchors = []
            w.writerow([code, name, rec["SIGNGU_CD_"], inside, bn, bl,
                        (bd if bd != "" else ""),
                        "Y" if (bd != "" and bd <= YEOKSEGWON_M) else "N",
                        "Y" if transfer_near else "N",
                        json.dumps(anchors, ensure_ascii=False), as_of or ""])

    print(f"\n원천 역 {raw_n} → 서울 행정구역 {seoul_n} + 인접 경기 {margin_n} (제외 {drop_n})")
    print(f"역-노선 레코드: {len(station_rows)}  (환승역 {sum(1 for s in station_rows if s['환승역']=='Y')})")
    print(f"상권 배정 성공: {sum(1 for s in station_rows if s['상권_코드'])}")
    print(f"상권 {len(t_recs)} 중 {NEAR_M}m 내 역 있는 상권: {n_with}  "
          f"(역세권 250m: {sum(1 for c,_,_ in t_recs if near[c] and sorted(near[c])[0][0] <= YEOKSEGWON_M)})")
    print(f"→ {st_path}")
    print(f"→ {sg_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
